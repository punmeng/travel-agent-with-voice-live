import json
import sys
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.stdin.reconfigure(encoding="utf-8")
sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")

CITY = ["city_id", "city_name", "alias", "country_code", "airport_iata", "airport_name", "timezone", "source", "updated_at"]
ROUTES = ["origin_iata", "destination_iata", "marketing_carrier", "operating_carrier", "valid_from", "valid_to", "source", "updated_at"]
HEADERS = {"CityAirportMapping": CITY, "RouteReferences": ROUTES}


def create(path):
    data = json.load(sys.stdin)
    book = Workbook()
    book.remove(book.active)
    for name, columns in HEADERS.items():
        sheet = book.create_sheet(name)
        sheet.append(columns)
        rows = data["rows"] if name == "CityAirportMapping" else data["routeReferences"]
        for row in rows:
            sheet.append([row[column] for column in columns])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column, width in zip("ABCDEFGHI", [28, 18, 24, 17, 17, 22, 26, 72, 18]):
            sheet.column_dimensions[column].width = width
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Arial", size=11)
                cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="263959")
        if rows:
            table = Table(displayName=name, ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
    guide = book.create_sheet("README")
    for line in [
        "城市／機場地理對照範本；不代表任何航空公司的實際營運航網。",
        "一列代表 city_id + alias + airport_iata；同一城市每個 alias 必須列出相同完整機場集合。",
        "所有欄位為文字；日期填 YYYY-MM-DD，不用公式、巨集、外部連結或 Excel 日期數字。",
        "城市別名只按 NFKC／大小寫／空白正規化作確切比對，不自動模糊猜城市。",
        "同名城市可在不同國家存在；country_code 用兩碼大寫，工具會回 ambiguous 並詢問國家。",
        "機場必須已在 data/network.json，名稱／IANA 時區／國家必須與既有 inventory 一致。",
        "本版每城市最多兩座候選機場，保留既有搜尋 fan-out 上限。",
        "source 保留可追溯來源；預設列是人工整理參考，不是假裝官方資料。",
        "RouteReferences 可空白；填入的資料僅供獨立參考，不建立可售資格或更動合成航班生成器。",
        "路線參考不驗證航空公司的實際營運資格。",
        "先執行 node scripts\\import-city-airports.js <檔案.xlsx> 檢查，再加 --apply 明確套用。",
        "套用是完整地理對照快照替換，不更動原始 network／班表／價格；服務受控重啟後載入。",
        "範本不含公式、巨集、樣本航空公司航線或個人資料；整份 Excel 不送往 LLM／向量庫。",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    for row in guide:
        row[0].font = Font(name="Arial", size=11)
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row[0].row].height = 32
    book.save(path)


def read(path):
    if path.suffix.lower() != ".xlsx" or path.stat().st_size > 2_000_000:
        raise ValueError("Only .xlsx up to 2 MB is accepted")
    with ZipFile(path) as archive:
        items = archive.infolist()
        if len(items) > 1000 or sum(item.file_size for item in items) > 20_000_000:
            raise ValueError("Workbook expanded size limit exceeded")
        for item in items:
            name = item.filename.lower()
            if any(part in name for part in ["vbaproject", "externallinks/", "embeddings/", "connections.xml", "activex/"]):
                raise ValueError("Macros, external data and embedded objects are not accepted")
            if name.endswith(".rels"):
                if any(element.get("TargetMode") == "External" for element in ElementTree.fromstring(archive.read(item))):
                    raise ValueError("External relationships are not accepted")
    book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if "CityAirportMapping" not in book.sheetnames or any(name not in [*HEADERS, "README"] for name in book.sheetnames):
            raise ValueError("Expected CityAirportMapping, optional RouteReferences and README sheets only")
        result = {"schemaVersion": 1, "rows": [], "routeReferences": []}
        for sheet in book:
            if sheet.max_row > 10001 or sheet.max_column > (1 if sheet.title == "README" else len(HEADERS[sheet.title])):
                raise ValueError(f"{sheet.title}: row/column limit exceeded")
            if sheet.sheet_state != "visible":
                raise ValueError("Hidden sheets are not accepted")
            rows = []
            for cells in sheet:
                values = []
                for cell in cells:
                    if cell.data_type in ["f", "e"] or (cell.value is not None and not isinstance(cell.value, str)):
                        raise ValueError(f"{sheet.title}!{cell.coordinate}: only plain text, no formulas/errors")
                    if cell.value and len(cell.value) > 1000:
                        raise ValueError(f"{sheet.title}!{cell.coordinate}: text too long")
                    values.append(cell.value.strip() if cell.value else "")
                rows.append(values)
            if sheet.title == "README":
                continue
            columns = HEADERS[sheet.title]
            if not rows or rows[0] != columns:
                raise ValueError(f"{sheet.title}: exact documented column names/order required")
            target = "rows" if sheet.title == "CityAirportMapping" else "routeReferences"
            for number, values in enumerate(rows[1:], 2):
                if not any(values):
                    continue
                if len(values) != len(columns) or not all(values):
                    raise ValueError(f"{sheet.title}!row {number}: all columns are required")
                result[target].append(dict(zip(columns, values)))
        return result
    finally:
        book.close()


if __name__ == "__main__":
    try:
        command, filename = sys.argv[1:]
        path = Path(filename)
        if command == "create":
            create(path)
            print(json.dumps({"created": True}))
        elif command == "read":
            print(json.dumps(read(path), ensure_ascii=True))
        else:
            raise ValueError("Expected create or read")
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
